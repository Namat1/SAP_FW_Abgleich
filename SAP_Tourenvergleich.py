import base64
import html
import io
import json
import hashlib
from typing import Dict, List, Optional, Set, Tuple

import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------------------
# Grundeinstellungen
# ---------------------------------------------------------------------------

DAY_NAMES = {
    1: "Montag",
    2: "Dienstag",
    3: "Mittwoch",
    4: "Donnerstag",
    5: "Freitag",
    6: "Samstag",
}

DAY_SHORT = {
    1: "Mo",
    2: "Di",
    3: "Mi",
    4: "Do",
    5: "Fr",
    6: "Sa",
}

# Fallback-Positionen, falls Spaltenüberschriften nicht erkannt werden.
# Tourenplanung: A CSB, B SAP, C Name, D Straße, E PLZ, F Ort, G Mo ... L Sa
DAY_COLUMNS_TOUR = {
    1: 6,
    2: 7,
    3: 8,
    4: 9,
    5: 10,
    6: 11,
}

SAP_COL_INDEX = 0
SAP_DAY_COL_INDEX = 6
TOUR_SAP_COL_INDEX = 1

# Nur diese Blätter der Tourenplanung bilden den Soll-Stand.
TOUR_SHEET_CANDIDATES = [
    "DIREKT",
    "MK",
    "HUPA_NMS",
    "HUPA_MALCHOW",
]

DAY_COLUMN_CANDIDATES = {
    1: ["mo", "montag"],
    2: ["die", "di", "dienstag"],
    3: ["mitt", "mit", "mi", "mittwoch"],
    4: ["don", "do", "donnerstag"],
    5: ["fr", "frei", "freitag"],
    6: ["sam", "sa", "samstag"],
}


# ---------------------------------------------------------------------------
# Helfer
# ---------------------------------------------------------------------------


def value_to_clean_text(value) -> str:
    if value is None or pd.isna(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_sap_series(series: pd.Series) -> pd.Series:
    """Normalisiert SAP-Nummern. Aus 12345.0 wird 12345."""
    if series.empty:
        return series.astype(str)

    result = series.copy()
    numeric = pd.to_numeric(result, errors="coerce")
    is_int = numeric.notna() & (numeric == numeric.round())

    out = result.astype(str)
    out = out.where(~is_int, numeric.where(is_int).astype("Int64").astype(str))
    out = out.str.strip()
    return out.replace({"nan": "", "<NA>": "", "None": ""})


def normalize_header_name(value) -> str:
    text = "" if value is None or pd.isna(value) else str(value)
    text = text.strip().lower()
    text = (
        text.replace("ä", "ae")
        .replace("ö", "oe")
        .replace("ü", "ue")
        .replace("ß", "ss")
    )
    return "".join(ch for ch in text if ch.isalnum())


def normalized_candidates(values: List[str]) -> List[str]:
    return [normalize_header_name(value) for value in values]


def pick_first_matching_column(columns: List[str], candidates: List[str]) -> Optional[str]:
    candidate_set = set(candidates)
    for column in columns:
        if normalize_header_name(column) in candidate_set:
            return column
    return None


def pick_column_by_name_or_position(
    columns: List[str],
    candidates: List[str],
    fallback_index: Optional[int] = None,
) -> Optional[str]:
    found = pick_first_matching_column(columns, normalized_candidates(candidates))
    if found is not None:
        return found
    if fallback_index is not None and len(columns) > fallback_index:
        return columns[fallback_index]
    return None


def make_unique_columns(raw_columns: List[object]) -> List[str]:
    result: List[str] = []
    seen: Dict[str, int] = {}
    for index, value in enumerate(raw_columns, start=1):
        name = value_to_clean_text(value) or f"Spalte_{index}"
        count = seen.get(name, 0) + 1
        seen[name] = count
        if count > 1:
            name = f"{name}_{count}"
        result.append(name)
    return result


def read_excel_with_detected_header(excel: pd.ExcelFile, sheet_name: str) -> pd.DataFrame:
    """Erkennt die Kopfzeile anhand von SAP und mindestens zwei Tages-Spalten."""
    raw = pd.read_excel(excel, sheet_name=sheet_name, header=None, dtype=object)
    if raw.empty:
        return pd.DataFrame()

    header_row: Optional[int] = None
    max_scan_rows = min(len(raw), 25)
    day_names_flat = {
        normalize_header_name(candidate)
        for values in DAY_COLUMN_CANDIDATES.values()
        for candidate in values
    }

    for row_index in range(max_scan_rows):
        values = [normalize_header_name(value) for value in raw.iloc[row_index].tolist()]
        value_set = set(values)
        has_sap = "sap" in value_set or "sapnummer" in value_set or "sapnr" in value_set
        day_hits = sum(1 for value in values if value in day_names_flat)
        if has_sap and day_hits >= 2:
            header_row = row_index
            break

    if header_row is None:
        return pd.read_excel(excel, sheet_name=sheet_name, header=0, dtype=object)

    df = raw.iloc[header_row + 1:].copy()
    df.columns = make_unique_columns(raw.iloc[header_row].tolist())
    return df.dropna(how="all").reset_index(drop=True)


def select_tour_sheet_names(excel: pd.ExcelFile) -> Tuple[List[str], List[str]]:
    available = {normalize_header_name(name): name for name in excel.sheet_names}
    selected: List[str] = []
    missing: List[str] = []

    for expected in TOUR_SHEET_CANDIDATES:
        real_name = available.get(normalize_header_name(expected))
        if real_name:
            selected.append(real_name)
        else:
            missing.append(expected)

    if selected:
        return selected, missing

    # Falls Blattnamen stark abweichen: die ersten vier Blätter als Fallback.
    return excel.sheet_names[:4], TOUR_SHEET_CANDIDATES


def day_value_is_set(value) -> bool:
    if value is None or pd.isna(value):
        return False
    text = str(value).strip()
    if not text or text.lower() in {"nan", "none", "<na>", "-", "--"}:
        return False
    number = pd.to_numeric(pd.Series([value]), errors="coerce").iloc[0]
    if pd.notna(number) and float(number) == 0:
        return False
    return True


def normalize_day_code_series(series: pd.Series) -> pd.Series:
    numeric = pd.to_numeric(series, errors="coerce")
    text = series.astype(str).map(normalize_header_name)
    text_map = {
        "1": 1, "mo": 1, "montag": 1,
        "2": 2, "di": 2, "die": 2, "dienstag": 2,
        "3": 3, "mi": 3, "mitt": 3, "mit": 3, "mittwoch": 3,
        "4": 4, "do": 4, "don": 4, "donnerstag": 4,
        "5": 5, "fr": 5, "frei": 5, "freitag": 5,
        "6": 6, "sa": 6, "sam": 6, "samstag": 6,
    }
    mapped = text.map(text_map)
    return numeric.where(numeric.notna(), mapped)


def days_to_text(days: Set[int] | List[int]) -> str:
    return ", ".join(DAY_SHORT[d] for d in sorted(days))


def uploaded_file_fingerprint(uploaded_file) -> str:
    """Eindeutiger Fingerabdruck des aktuell hochgeladenen Dateiinhalts."""
    if uploaded_file is None:
        return ""
    return hashlib.sha256(uploaded_file.getvalue()).hexdigest()


def merge_customer_info(base: Dict[str, Dict[str, str]], sap: str, info: Dict[str, str]) -> None:
    target = base.setdefault(sap, {"name": "", "strasse": "", "ort": ""})
    for key in ["name", "strasse", "ort"]:
        if not target.get(key) and info.get(key):
            target[key] = info[key]


# ---------------------------------------------------------------------------
# Dateien lesen
# ---------------------------------------------------------------------------


def read_sap_file(uploaded_file) -> Tuple[Dict[str, Set[int]], Set[str], str, int]:
    """SAP Ist-Stand: Fallback A = SAP, G = Liefertag."""
    excel = pd.ExcelFile(uploaded_file)
    sheet_name = excel.sheet_names[0]
    df = read_excel_with_detected_header(excel, sheet_name)

    if df.empty:
        return {}, set(), sheet_name, 0

    columns = list(df.columns)
    sap_column = pick_column_by_name_or_position(
        columns,
        ["SAP", "SAP Nummer", "SAP-Nr", "SAP Nr", "Kundennummer", "Kunden Nummer"],
        SAP_COL_INDEX,
    )
    day_column = pick_column_by_name_or_position(
        columns,
        ["Liefertag", "Liefer Tag", "LT", "Tag", "Liefertag Code", "Liefertagcode"],
        SAP_DAY_COL_INDEX,
    )

    if sap_column is None or day_column is None:
        return {}, set(), sheet_name, 0

    work = df[[sap_column, day_column]].copy()
    work.columns = ["sap", "tag"]
    work["sap"] = normalize_sap_series(work["sap"])

    # Kundenbestand unabhängig vom Liefertag erfassen. So kann eindeutig
    # unterschieden werden: Kunde fehlt komplett in SAP vs. Liefertag weicht ab.
    sap_customers: Set[str] = set(work.loc[work["sap"].ne(""), "sap"].tolist())

    work["tag_num"] = normalize_day_code_series(work["tag"])

    mask = (
        work["sap"].ne("")
        & work["tag_num"].notna()
        & work["tag_num"].between(1, 6, inclusive="both")
    )
    filtered = work.loc[mask, ["sap", "tag_num"]].copy()
    filtered["tag_int"] = filtered["tag_num"].astype(int)

    days_by_sap: Dict[str, Set[int]] = filtered.groupby("sap")["tag_int"].agg(set).to_dict()
    return days_by_sap, sap_customers, sheet_name, len(filtered)


def read_tourenplanung(
    uploaded_file,
) -> Tuple[pd.DataFrame, List[str], List[str], Dict[str, Dict[str, str]]]:
    """Tourenplanung = verbindlicher Soll-Stand."""
    excel = pd.ExcelFile(uploaded_file)
    sheet_names, missing_sheet_names = select_tour_sheet_names(excel)

    frames: List[pd.DataFrame] = []
    customer_info: Dict[str, Dict[str, str]] = {}

    for sheet_name in sheet_names:
        df = read_excel_with_detected_header(excel, sheet_name)
        if df.empty:
            continue

        columns = list(df.columns)
        sap_column = pick_column_by_name_or_position(
            columns,
            ["SAP", "SAP Nummer", "SAP-Nr", "SAP Nr", "Kundennummer", "Kunden Nummer"],
            TOUR_SAP_COL_INDEX,
        )
        if sap_column is None:
            continue

        name_column = pick_column_by_name_or_position(
            columns,
            ["Name", "Kundenname", "Marktname", "Kunde", "Bezeichnung", "Filialname"],
            2,
        )
        strasse_column = pick_column_by_name_or_position(
            columns,
            ["Strasse", "Straße", "Str", "Anschrift", "Adresse", "Strassenname", "Straßenname", "Strasse Hausnummer", "Straße Hausnummer"],
            3,
        )
        plz_column = pick_column_by_name_or_position(columns, ["Plz", "PLZ", "Postleitzahl"], 4)
        ort_column = pick_column_by_name_or_position(columns, ["Ort", "Stadt", "Plz Ort", "PLZ Ort", "Ortname"], 5)

        rename_map = {sap_column: "sap"}
        for day_num, col_index in DAY_COLUMNS_TOUR.items():
            day_column = pick_column_by_name_or_position(
                columns,
                DAY_COLUMN_CANDIDATES[day_num],
                col_index,
            )
            if day_column and day_column != sap_column:
                rename_map[day_column] = f"tag_{day_num}"

        if name_column and name_column != sap_column:
            rename_map[name_column] = "name"
        if strasse_column and strasse_column != sap_column:
            rename_map[strasse_column] = "strasse"
        if ort_column and ort_column != sap_column:
            rename_map[ort_column] = "ort"
        if plz_column and plz_column != sap_column:
            rename_map[plz_column] = "plz"

        work = df.rename(columns=rename_map).copy()
        work["sap"] = normalize_sap_series(work["sap"])
        work = work[work["sap"].ne("")].copy()
        if work.empty:
            continue

        for column in ["name", "strasse", "ort", "plz"]:
            if column not in work.columns:
                work[column] = ""

        info_df = work[["sap", "name", "strasse", "ort", "plz"]].copy()
        for column in ["name", "strasse", "ort", "plz"]:
            info_df[column] = info_df[column].map(value_to_clean_text)
        info_df["ort_kombi"] = info_df.apply(
            lambda row: " ".join(v for v in [row["plz"], row["ort"]] if v).strip(),
            axis=1,
        )

        for _, row in info_df.iterrows():
            merge_customer_info(
                customer_info,
                row["sap"],
                {
                    "name": row["name"],
                    "strasse": row["strasse"],
                    "ort": row["ort_kombi"] or row["ort"],
                },
            )

        day_value_columns = [
            f"tag_{d}" for d in DAY_COLUMNS_TOUR if f"tag_{d}" in work.columns
        ]
        if not day_value_columns:
            continue

        work["blatt"] = sheet_name
        long = work.melt(
            id_vars=["sap", "blatt"],
            value_vars=day_value_columns,
            var_name="tag_col",
            value_name="wert",
        )
        long["tag_num"] = long["tag_col"].str.replace("tag_", "", regex=False).astype(int)
        long["wert_gesetzt"] = long["wert"].map(day_value_is_set)
        long = long[long["sap"].ne("") & long["wert_gesetzt"]]
        frames.append(long[["sap", "blatt", "tag_num"]])

    if not frames:
        empty = pd.DataFrame(columns=["sap", "blatt", "tag_num"])
        return empty, sheet_names, missing_sheet_names, customer_info

    return pd.concat(frames, ignore_index=True), sheet_names, missing_sheet_names, customer_info


# ---------------------------------------------------------------------------
# Vergleich + Gesamtübersicht
# ---------------------------------------------------------------------------

RESULT_COLUMNS = [
    "Blatt",
    "SAP Nummer",
    "Name",
    "Straße",
    "Ort",
    "Tourenplanung (Soll)",
    "SAP (Ist)",
    "SAP-Abweichung",
]

OVERVIEW_COLUMNS = [
    "Blatt",
    "SAP Nummer",
    "Name",
    "Straße",
    "Ort",
    "Tourenplanung (Soll)",
    "SAP (Ist)",
    "Fehlt in SAP",
    "Zusätzlich in SAP",
    "Status",
    "_tour_days",
    "_sap_days",
]


def build_sap_differences(
    tour_df: pd.DataFrame,
    days_by_sap: Dict[str, Set[int]],
    sap_customers: Set[str],
    customer_info: Dict[str, Dict[str, str]],
) -> pd.DataFrame:
    """Die Tourenplanung ist der Soll-Stand; SAP-Abweichungen werden vollständig gezeigt."""
    if tour_df.empty:
        return pd.DataFrame(columns=RESULT_COLUMNS)

    expected_by_sap: Dict[str, Set[int]] = tour_df.groupby("sap")["tag_num"].agg(set).to_dict()
    sheets_by_sap: Dict[str, str] = tour_df.groupby("sap")["blatt"].agg(
        lambda x: ", ".join(sorted(set(map(str, x))))
    ).to_dict()

    rows: List[dict] = []
    for sap, expected_days in expected_by_sap.items():
        info = customer_info.get(sap, {})
        if sap not in sap_customers:
            rows.append({
                "Blatt": sheets_by_sap.get(sap, ""),
                "SAP Nummer": sap,
                "Name": info.get("name", ""),
                "Straße": info.get("strasse", ""),
                "Ort": info.get("ort", ""),
                "Tourenplanung (Soll)": days_to_text(expected_days) or "–",
                "SAP (Ist)": "Kunde fehlt",
                "SAP-Abweichung": "Kunde fehlt in SAP",
                "_sort": int(sap) if str(sap).isdigit() else 9_999_999_999,
            })
            continue

        actual_days = set(days_by_sap.get(sap, set()))
        missing_in_sap = sorted(set(expected_days) - actual_days)
        extra_in_sap = sorted(actual_days - set(expected_days))

        if not missing_in_sap and not extra_in_sap:
            continue

        parts: List[str] = []
        if missing_in_sap:
            parts.append(f"Fehlt in SAP: {days_to_text(missing_in_sap)}")
        if extra_in_sap:
            parts.append(f"Zusätzlich in SAP: {days_to_text(extra_in_sap)}")

        rows.append({
            "Blatt": sheets_by_sap.get(sap, ""),
            "SAP Nummer": sap,
            "Name": info.get("name", ""),
            "Straße": info.get("strasse", ""),
            "Ort": info.get("ort", ""),
            "Tourenplanung (Soll)": days_to_text(expected_days) or "–",
            "SAP (Ist)": days_to_text(actual_days) or "–",
            "SAP-Abweichung": " | ".join(parts),
            "_sort": int(sap) if str(sap).isdigit() else 9_999_999_999,
        })

    if not rows:
        return pd.DataFrame(columns=RESULT_COLUMNS)

    result = pd.DataFrame(rows)
    result = result.sort_values(["Blatt", "_sort"]).reset_index(drop=True)
    return result[RESULT_COLUMNS]


def build_customer_overview(
    tour_df: pd.DataFrame,
    days_by_sap: Dict[str, Set[int]],
    sap_customers: Set[str],
    customer_info: Dict[str, Dict[str, str]],
) -> pd.DataFrame:
    """Eine Zeile je Kunde aus der Quelldatei mit Soll/Ist direkt nebeneinander."""
    if tour_df.empty:
        return pd.DataFrame(columns=OVERVIEW_COLUMNS)

    expected_by_sap: Dict[str, Set[int]] = tour_df.groupby("sap")["tag_num"].agg(set).to_dict()
    sheets_by_sap: Dict[str, str] = tour_df.groupby("sap")["blatt"].agg(
        lambda x: ", ".join(sorted(set(map(str, x))))
    ).to_dict()

    rows: List[dict] = []
    for sap, expected_days_raw in expected_by_sap.items():
        expected_days = set(expected_days_raw)
        exists = sap in sap_customers
        actual_days = set(days_by_sap.get(sap, set())) if exists else set()
        missing = expected_days - actual_days
        extra = actual_days - expected_days
        info = customer_info.get(sap, {})

        if not exists:
            status = "Kunde fehlt in SAP"
        elif missing or extra:
            status = "Abweichung"
        else:
            status = "OK"

        rows.append({
            "Blatt": sheets_by_sap.get(sap, ""),
            "SAP Nummer": str(sap),
            "Name": info.get("name", ""),
            "Straße": info.get("strasse", ""),
            "Ort": info.get("ort", ""),
            "Tourenplanung (Soll)": days_to_text(expected_days) or "–",
            "SAP (Ist)": days_to_text(actual_days) if exists and actual_days else ("–" if exists else "Kunde fehlt"),
            "Fehlt in SAP": days_to_text(missing) or "–",
            "Zusätzlich in SAP": days_to_text(extra) or "–",
            "Status": status,
            "_tour_days": expected_days,
            "_sap_days": actual_days,
            "_sort": int(sap) if str(sap).isdigit() else 9_999_999_999,
        })

    result = pd.DataFrame(rows)
    result = result.sort_values(["Blatt", "_sort"]).drop(columns=["_sort"]).reset_index(drop=True)
    return result[OVERVIEW_COLUMNS]


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------


def build_excel(overview: pd.DataFrame, differences: pd.DataFrame) -> bytes:
    output = io.BytesIO()
    overview_export = overview.drop(columns=["_tour_days", "_sap_days"], errors="ignore").copy()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        overview_export.to_excel(writer, index=False, sheet_name="Gesamtübersicht", na_rep="")
        differences.to_excel(writer, index=False, sheet_name="SAP Abweichungen", na_rep="")

        header_fill = PatternFill(start_color="FF2F3A4A", end_color="FF2F3A4A", fill_type="solid")
        ok_fill = PatternFill(start_color="FFEAF7F0", end_color="FFEAF7F0", fill_type="solid")
        diff_fill = PatternFill(start_color="FFFFF7D6", end_color="FFFFF7D6", fill_type="solid")
        missing_fill = PatternFill(start_color="FFFEE2E2", end_color="FFFEE2E2", fill_type="solid")
        thin = Side(style="thin", color="FFD7DEE8")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        width_hints = {
            "Blatt": 18, "SAP Nummer": 13, "Name": 32, "Straße": 26, "Ort": 26,
            "Tourenplanung (Soll)": 22, "SAP (Ist)": 22, "Fehlt in SAP": 20,
            "Zusätzlich in SAP": 22, "Status": 20, "SAP-Abweichung": 38,
        }

        for sheet_name, df in [("Gesamtübersicht", overview_export), ("SAP Abweichungen", differences)]:
            ws = writer.sheets[sheet_name]
            columns = list(df.columns)

            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
                cell.alignment = Alignment(vertical="center")
                cell.border = border

            for row_idx in range(2, len(df) + 2):
                row_fill = diff_fill
                if sheet_name == "Gesamtübersicht" and "Status" in columns:
                    status = str(ws.cell(row=row_idx, column=columns.index("Status") + 1).value or "")
                    if status == "OK":
                        row_fill = ok_fill
                    elif status == "Kunde fehlt in SAP":
                        row_fill = missing_fill
                for col_idx in range(1, len(columns) + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.font = Font(name="Calibri", size=10)
                    cell.alignment = Alignment(vertical="center", wrap_text=False)
                    cell.border = border
                    cell.fill = row_fill

            for idx, column in enumerate(columns, start=1):
                ws.column_dimensions[get_column_letter(idx)].width = width_hints.get(column, 20)

            ws.freeze_panes = "A2"
            if columns:
                ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(df) + 1}"
            ws.sheet_view.showGridLines = False
            ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE

    return output.getvalue()


# ---------------------------------------------------------------------------
# HTML: alle Kunden, Tour/Quelldatei und SAP nebeneinander
# ---------------------------------------------------------------------------


def build_html_report(
    overview: pd.DataFrame,
    differences: pd.DataFrame,
    sap_sheet: str,
    tour_sheets: List[str],
    excel_bytes: bytes,
) -> bytes:
    excel_b64 = base64.b64encode(excel_bytes).decode("ascii")
    total_count = len(overview)
    diff_count = len(differences)
    ok_count = max(total_count - diff_count, 0)
    missing_customer_count = int((overview["Status"] == "Kunde fehlt in SAP").sum()) if not overview.empty else 0
    extra_customer_count = 0
    if not overview.empty:
        extra_customer_count = sum(bool(set(row["_sap_days"]) - set(row["_tour_days"])) for _, row in overview.iterrows())

    def day_badges(days: Set[int], missing: Set[int] | None = None, extra: Set[int] | None = None, mode: str = "tour") -> str:
        missing = missing or set()
        extra = extra or set()
        if not days:
            return "<span class='none'>–</span>"
        parts: List[str] = []
        for day in sorted(days):
            label = DAY_SHORT.get(day, str(day))
            if day in missing:
                parts.append(f"<span class='day day-missing' title='Fehlt in SAP'>{label}</span>")
            elif day in extra:
                parts.append(f"<span class='day day-extra' title='Zusätzlich in SAP'>+{label}</span>")
            elif mode == "sap":
                parts.append(f"<span class='day day-sap'>{label}</span>")
            else:
                parts.append(f"<span class='day day-tour'>{label}</span>")
        return "".join(parts)

    row_html: List[str] = []
    for _, row in overview.iterrows():
        tour_days = set(row.get("_tour_days", set()))
        sap_days = set(row.get("_sap_days", set()))
        missing = tour_days - sap_days
        extra = sap_days - tour_days
        status = str(row.get("Status", ""))
        filter_status = "ok" if status == "OK" else "diff"
        status_cls = "badge-ok" if status == "OK" else ("badge-missing" if status == "Kunde fehlt in SAP" else "badge-diff")
        row_cls = "row-ok" if status == "OK" else "row-diff"

        tour_html = day_badges(tour_days, missing=missing, mode="tour")
        sap_html = day_badges(sap_days, extra=extra, mode="sap")
        missing_html = day_badges(missing, missing=missing) if missing else "<span class='none'>–</span>"
        extra_html = day_badges(extra, extra=extra) if extra else "<span class='none'>–</span>"

        row_html.append(
            f"<tr class='{row_cls}' data-status='{filter_status}'>"
            f"<td><span class='area'>{html.escape(str(row.get('Blatt', '')))}</span></td>"
            f"<td class='mono'>{html.escape(str(row.get('SAP Nummer', '')))}</td>"
            f"<td class='name'>{html.escape(str(row.get('Name', '')))}</td>"
            f"<td>{html.escape(str(row.get('Straße', '')))}</td>"
            f"<td>{html.escape(str(row.get('Ort', '')))}</td>"
            f"<td class='days-cell'><div class='source-label tour-label'>QUELLE</div><div class='days'>{tour_html}</div></td>"
            f"<td class='days-cell'><div class='source-label sap-label'>SAP</div><div class='days'>{sap_html}</div></td>"
            f"<td class='days-cell'><div class='days'>{missing_html}</div></td>"
            f"<td class='days-cell'><div class='days'>{extra_html}</div></td>"
            f"<td><span class='status-badge {status_cls}'>{html.escape(status)}</span></td>"
            "</tr>"
        )

    sheet_chips = "".join(f"<span class='chip'>{html.escape(str(s))}</span>" for s in tour_sheets)
    table = f"""
    <div class='table-wrap'>
      <table id='resultTable'>
        <thead><tr>
          <th>Blatt</th><th>SAP Nummer</th><th>Kunde</th><th>Straße</th><th>Ort</th>
          <th class='tour-head'>Quelldatei / Soll</th><th class='sap-head'>SAP / Ist</th>
          <th>Fehlt in SAP</th><th>Zusätzlich in SAP</th><th>Status</th>
        </tr></thead>
        <tbody>{''.join(row_html)}</tbody>
      </table>
    </div>
    """


    graph_nodes: List[dict] = [{
        "id": "root", "type": "root", "label": "FW SAP", "status": "root",
        "name": "FW SAP – Quelldatei Abgleich", "group": "", "sap": "", "source": "", "sap_days": "",
        "missing": "", "extra": "", "detail_status": "Gesamtauswertung", "parent": ""
    }]
    graph_edges: List[dict] = []
    graph_groups: Dict[str, str] = {}

    for graph_index, (_, graph_row) in enumerate(overview.iterrows()):
        graph_tour_days = set(graph_row.get("_tour_days", set()))
        graph_sap_days = set(graph_row.get("_sap_days", set()))
        graph_missing = graph_tour_days - graph_sap_days
        graph_extra = graph_sap_days - graph_tour_days
        graph_status_text = str(graph_row.get("Status", ""))
        graph_status = "missing" if (graph_status_text == "Kunde fehlt in SAP" or graph_missing) else ("extra" if graph_extra else "ok")

        graph_group_names = [part.strip() for part in str(graph_row.get("Blatt", "")).split(",") if part.strip()]
        if not graph_group_names:
            graph_group_names = ["Ohne Blatt"]
        graph_group_ids: List[str] = []
        for graph_group_name in graph_group_names:
            if graph_group_name not in graph_groups:
                graph_group_id = f"group_{len(graph_groups)}"
                graph_groups[graph_group_name] = graph_group_id
                graph_nodes.append({
                    "id": graph_group_id, "type": "group", "label": graph_group_name, "status": "group",
                    "name": graph_group_name, "group": graph_group_name, "sap": "", "source": "",
                    "sap_days": "", "missing": "", "extra": "", "detail_status": "Gruppe", "parent": "root"
                })
                graph_edges.append({"source": "root", "target": graph_group_id})
            graph_group_ids.append(graph_groups[graph_group_name])

        graph_customer_id = f"customer_{graph_index}"
        graph_customer_name = str(graph_row.get("Name", "")).strip() or str(graph_row.get("SAP Nummer", "")).strip() or "Kunde"
        graph_nodes.append({
            "id": graph_customer_id, "type": "customer", "label": graph_customer_name, "status": graph_status,
            "name": graph_customer_name, "group": ", ".join(graph_group_names),
            "sap": str(graph_row.get("SAP Nummer", "")),
            "source": days_to_text(graph_tour_days) or "–", "sap_days": days_to_text(graph_sap_days) or "–",
            "missing": days_to_text(graph_missing) or "–", "extra": days_to_text(graph_extra) or "–",
            "detail_status": graph_status_text, "parent": graph_group_ids[0]
        })
        for graph_group_id in graph_group_ids:
            graph_edges.append({"source": graph_group_id, "target": graph_customer_id})

    graph_json = json.dumps({"nodes": graph_nodes, "edges": graph_edges}, ensure_ascii=False)
    graph_json = graph_json.replace("<", "\u003c").replace(">", "\u003e").replace("&", "\u0026")

    report = f"""<!doctype html>
<html lang='de'>
<head>
<meta charset='utf-8'>
<meta name='viewport' content='width=device-width,initial-scale=1'>
<title>FW SAP – Quelldatei Abgleich</title>
<style>
:root {{
  --bg:#f5f6f8; --card:#fff; --text:#20242d; --muted:#667085; --line:#e2e5ea;
  --tour:#5b45a6; --tour-bg:#eeeafd; --sap:#176a47; --sap-bg:#e8f6ee;
  --bad:#b42318; --bad-bg:#feecea; --extra:#9a5b00; --extra-bg:#fff0cf; --extra-border:#f59e0b;
  --warn:#8a5a00; --warn-bg:#fff7dd;
}}
* {{ box-sizing:border-box; }}
body {{ margin:0; background:var(--bg); color:var(--text); font-family:Inter,Segoe UI,Arial,sans-serif; }}
.page {{ max-width:1780px; margin:0 auto; padding:30px 26px 48px; }}
.header {{ display:flex; justify-content:space-between; align-items:flex-start; gap:20px; margin-bottom:20px; }}
.eyebrow {{ font-size:12px; font-weight:900; letter-spacing:.10em; color:var(--tour); text-transform:uppercase; }}
h1 {{ margin:7px 0 7px; font-size:32px; line-height:1.1; }}
.subtitle {{ color:var(--muted); max-width:920px; line-height:1.5; }}
.download {{ text-decoration:none; color:#fff; background:#272b35; border-radius:12px; padding:13px 17px; font-weight:800; white-space:nowrap; }}
.grid {{ display:grid; grid-template-columns:repeat(4,minmax(0,1fr)); gap:12px; margin:18px 0; }}
.card {{ background:var(--card); border:1px solid var(--line); border-radius:15px; padding:15px 17px; box-shadow:0 3px 12px rgba(20,30,50,.035); }}
.metric-label {{ color:var(--muted); font-size:12px; font-weight:750; }}
.metric-value {{ font-size:28px; font-weight:900; margin:4px 0; }}
.metric-note {{ color:var(--muted); font-size:11px; }}
.info {{ display:flex; flex-wrap:wrap; gap:8px; margin-top:9px; }}
.chip {{ background:#f0eef9; color:#443589; border-radius:999px; padding:6px 10px; font-size:12px; font-weight:750; }}
.section-head {{ display:flex; justify-content:space-between; align-items:flex-end; gap:16px; margin:28px 0 12px; }}
.section-title {{ margin:0; font-size:21px; }}
.legend {{ display:flex; flex-wrap:wrap; gap:9px; color:var(--muted); font-size:12px; }}
.legend-item {{ display:inline-flex; align-items:center; gap:5px; }}
.legend-dot {{ width:10px; height:10px; border-radius:50%; }}
.legend-tour {{ background:var(--tour); }} .legend-sap {{ background:var(--sap); }} .legend-missing {{ background:var(--bad); }} .legend-extra {{ background:var(--extra-border); }}
.toolbar {{ display:flex; justify-content:space-between; align-items:center; gap:12px; margin:12px 0 10px; flex-wrap:wrap; }}
.search {{ flex:1 1 420px; max-width:680px; border:1px solid var(--line); background:#fff; border-radius:12px; padding:12px 14px; font-size:14px; outline:none; }}
.filters {{ display:flex; gap:7px; flex-wrap:wrap; }}
.filter-btn {{ border:1px solid var(--line); background:#fff; color:var(--text); border-radius:999px; padding:9px 13px; font-weight:800; cursor:pointer; }}
.filter-btn.active {{ background:#272b35; border-color:#272b35; color:#fff; }}
.result-count {{ color:var(--muted); font-size:13px; font-weight:800; }}
.table-wrap {{ overflow:visible; max-height:none; background:var(--card); border:1px solid var(--line); border-radius:16px; box-shadow:0 3px 12px rgba(20,30,50,.04); }}
table {{ width:100%; border-collapse:separate; border-spacing:0; table-layout:auto; }}
th {{ position:sticky; top:0; z-index:2; background:#252a36; color:#fff; text-align:left; font-size:11px; padding:12px 9px; }}
th:first-child {{ border-top-left-radius:15px; }} th:last-child {{ border-top-right-radius:15px; }}
th.tour-head {{ background:#53419c; }} th.sap-head {{ background:#176a47; }}
td {{ padding:10px 9px; border-top:1px solid var(--line); font-size:12px; vertical-align:middle; background:#fff; }}
tr.row-diff td {{ background:#fffcf5; }} tbody tr:hover td {{ background:#faf9fe; }}
.name {{ font-weight:800; min-width:150px; }} .mono {{ font-variant-numeric:tabular-nums; font-family:ui-monospace,SFMono-Regular,Consolas,monospace; }}
.area {{ display:inline-block; background:#f0f2f5; border-radius:999px; padding:5px 8px; font-weight:800; }}
.days {{ display:flex; gap:4px; flex-wrap:wrap; align-items:center; }} .source-label {{ font-size:9px; font-weight:950; letter-spacing:.08em; margin-bottom:4px; }}
.tour-label {{ color:var(--tour); }} .sap-label {{ color:var(--sap); }}
.day {{ min-width:30px; height:27px; display:inline-flex; align-items:center; justify-content:center; border-radius:7px; font-weight:900; font-size:11px; border:1px solid transparent; }}
.day-tour {{ background:var(--tour-bg); color:#4d3a9b; border-color:#d9d1f4; }}
.day-sap {{ background:var(--sap-bg); color:var(--sap); border-color:#b7e2c9; }}
.day-missing {{ background:var(--bad-bg); color:var(--bad); border-color:#f5bbb5; }}
.day-extra {{ background:var(--extra-bg); color:var(--extra); border:2px solid var(--extra-border); font-weight:950; }}
.none {{ color:#a0a7b2; }}
.status-badge {{ display:inline-flex; border-radius:999px; padding:6px 9px; font-size:10px; font-weight:900; white-space:nowrap; }}
.badge-ok {{ background:var(--sap-bg); color:var(--sap); }} .badge-diff {{ background:var(--warn-bg); color:var(--warn); }} .badge-missing {{ background:var(--bad-bg); color:var(--bad); }}

.view-switch {{ display:flex; gap:8px; margin:14px 0 12px; flex-wrap:wrap; }}
.view-btn {{ border:1px solid var(--line); background:#fff; color:var(--text); border-radius:11px; padding:10px 15px; font-weight:850; cursor:pointer; }}
.view-btn.active {{ color:#fff; background:#272b35; border-color:#272b35; }}
.view-panel.hidden {{ display:none; }}
.graph-intro {{ display:flex; justify-content:space-between; gap:14px; align-items:center; margin:8px 0 12px; color:var(--muted); font-size:12px; flex-wrap:wrap; }}
.graph-shell {{ position:relative; width:100%; min-height:760px; border-radius:18px; overflow:hidden; background:radial-gradient(circle at 50% 50%,#202534 0,#141821 42%,#0d1016 100%); border:1px solid #2c3342; box-shadow:0 12px 30px rgba(13,16,22,.16); }}
#graphSvg {{ display:block; width:100%; height:760px; cursor:grab; user-select:none; touch-action:none; }}
#graphSvg.panning {{ cursor:grabbing; }}
.graph-edge {{ stroke:#697386; stroke-opacity:.30; stroke-width:1.1; }}
.graph-node circle {{ stroke:#0d1016; stroke-width:2; transition:filter .12s ease; }}
.graph-node:hover circle {{ filter:drop-shadow(0 0 7px rgba(255,255,255,.35)); }}
.graph-node.root circle {{ fill:#8b5cf6; }}
.graph-node.group circle {{ fill:#c4b5fd; }}
.graph-node.ok circle {{ fill:#34d399; }}
.graph-node.extra circle {{ fill:#f59e0b; }}
.graph-node.missing circle {{ fill:#ef4444; }}
.graph-label {{ fill:#e7eaf0; font-size:12px; font-weight:800; pointer-events:none; paint-order:stroke; stroke:#11141b; stroke-width:4px; stroke-linejoin:round; }}
.graph-customer-label {{ fill:#f4f5f7; font-size:10px; opacity:0; pointer-events:none; paint-order:stroke; stroke:#11141b; stroke-width:3px; transition:opacity .12s; }}
.graph-node:hover .graph-customer-label {{ opacity:1; }}
.graph-detail {{ position:absolute; right:16px; top:16px; width:min(330px,calc(100% - 32px)); background:rgba(19,23,32,.94); color:#f6f7f9; border:1px solid #353d4d; border-radius:14px; padding:14px 15px; backdrop-filter:blur(10px); box-shadow:0 10px 28px rgba(0,0,0,.25); }}
.graph-detail-title {{ font-size:16px; font-weight:900; margin-bottom:8px; }}
.graph-detail-row {{ display:grid; grid-template-columns:92px 1fr; gap:8px; font-size:11px; line-height:1.45; padding:3px 0; }}
.graph-detail-key {{ color:#9ca6b6; font-weight:800; }}
.graph-help {{ position:absolute; left:14px; bottom:12px; color:#aab2c0; font-size:10px; background:rgba(13,16,22,.72); padding:7px 9px; border-radius:8px; }}
.graph-legend {{ display:flex; flex-wrap:wrap; gap:10px; }}
.graph-legend span {{ display:inline-flex; align-items:center; gap:5px; }}
.graph-dot {{ width:9px; height:9px; border-radius:50%; display:inline-block; }}
.graph-dot.ok {{ background:#34d399; }} .graph-dot.extra {{ background:#f59e0b; }} .graph-dot.missing {{ background:#ef4444; }} .graph-dot.group {{ background:#c4b5fd; }}

.footer {{ color:var(--muted); font-size:12px; margin-top:24px; text-align:center; }}
@media(max-width:1100px) {{ .grid{{grid-template-columns:repeat(2,1fr)}} th,td{{font-size:10px;padding:8px 6px}} .day{{min-width:26px;height:24px;font-size:10px}} }}
@media(max-width:760px) {{ .page{{padding:18px 10px 30px}} .header{{flex-direction:column}} .download{{width:100%;text-align:center}} .grid{{grid-template-columns:1fr}} .section-head{{flex-direction:column;align-items:flex-start}} }}
</style>
</head>
<body>
<div class='page'>
  <div class='header'>
    <div>
      <div class='eyebrow'>FW SAP · Quelldatei Abgleich</div>
      <h1>FW SAP – Quelldatei Abgleich</h1>
      <div class='subtitle'>Die Quelldatei ist der Soll-Stand. Alle Kunden werden vollständig dargestellt; Soll-Liefertage und SAP-Ist-Liefertage stehen direkt nebeneinander.</div>
    </div>
    <a class='download' href='data:application/vnd.openxmlformats-officedocument.spreadsheetml.sheet;base64,{excel_b64}' download='FW_SAP_Quelldatei_Abgleich.xlsx'>Excel herunterladen</a>
  </div>

  <div class='grid'>
    <div class='card'><div class='metric-label'>Geprüfte Kunden</div><div class='metric-value'>{total_count}</div><div class='metric-note'>aus der Quelldatei</div></div>
    <div class='card'><div class='metric-label'>Ohne Abweichung</div><div class='metric-value'>{ok_count}</div><div class='metric-note'>Soll und SAP identisch</div></div>
    <div class='card'><div class='metric-label'>Mit Abweichung</div><div class='metric-value'>{diff_count}</div><div class='metric-note'>fehlende oder zusätzliche SAP-Tage</div></div>
    <div class='card'><div class='metric-label'>Zusätzliche SAP-Tage</div><div class='metric-value'>{extra_customer_count}</div><div class='metric-note'>Kunden mit mindestens einem Extra-Tag</div></div>
  </div>

  <div class='card'>
    <div class='metric-label'>Geprüfte Datenbasis</div>
    <div style='margin-top:5px'><b>SAP-Blatt:</b> {html.escape(str(sap_sheet))}</div>
    <div class='info'>{sheet_chips}</div>
    <div style='margin-top:10px;color:var(--muted);font-size:12px'>Kunden, die nur in SAP stehen, werden weiterhin ignoriert. Fehlende Soll-Tage sind rot; zusätzliche SAP-Tage sind deutlich orange und mit <b>+</b> gekennzeichnet. Kunden, die komplett in SAP fehlen: <b>{missing_customer_count}</b>.</div>
  </div>

  <div class='section-head'>
    <h2 class='section-title'>Alle Kunden – Quelldatei und SAP nebeneinander</h2>
    <div class='legend'>
      <span class='legend-item'><i class='legend-dot legend-tour'></i> Quelldatei / Soll</span>
      <span class='legend-item'><i class='legend-dot legend-sap'></i> SAP passend</span>
      <span class='legend-item'><i class='legend-dot legend-missing'></i> fehlt in SAP</span>
      <span class='legend-item'><i class='legend-dot legend-extra'></i> + zusätzlich in SAP</span>
    </div>
  </div>

  <div class='view-switch'>
    <button id='tableViewBtn' class='view-btn active' onclick="showView('table')">Tabellenansicht</button>
    <button id='graphViewBtn' class='view-btn' onclick="showView('graph')">Stern-Graph</button>
  </div>

  <div id='tableView' class='view-panel'>
  <div class='toolbar'>
    <input id='searchInput' class='search' type='search' placeholder='SAP Nummer, Kunde, Straße, Ort, Blatt oder Liefertag suchen …' oninput='applyFilters()'>
    <div class='filters'>
      <button class='filter-btn active' onclick="setFilter('all',this)">Alle</button>
      <button class='filter-btn' onclick="setFilter('diff',this)">Nur Abweichungen</button>
      <button class='filter-btn' onclick="setFilter('ok',this)">Nur OK</button>
    </div>
    <span id='resultCount' class='result-count'></span>
  </div>

  {table}
  </div>

  <div id='graphView' class='view-panel hidden'>
    <div class='graph-intro'>
      <div><b>Obsidian-ähnlicher Stern-Graph:</b> Zentrum → Blatt → Kunde. Anklicken zeigt die Details.</div>
      <div class='graph-legend'>
        <span><i class='graph-dot group'></i> Blatt</span>
        <span><i class='graph-dot ok'></i> OK</span>
        <span><i class='graph-dot extra'></i> zusätzlicher SAP-Tag</span>
        <span><i class='graph-dot missing'></i> fehlt in SAP</span>
      </div>
    </div>
    <div class='graph-shell'>
      <svg id='graphSvg' role='img' aria-label='Stern-Graph der SAP-Auswertung'>
        <g id='graphViewport'><g id='graphEdges'></g><g id='graphNodes'></g></g>
      </svg>
      <div id='graphDetail' class='graph-detail'>
        <div class='graph-detail-title'>Stern-Graph</div>
        <div style='color:#aab2c0;font-size:11px;line-height:1.45'>Klicke einen Kunden an, um Soll-Tage, SAP-Tage und Abweichungen anzuzeigen.</div>
      </div>
      <div class='graph-help'>Mausrad: Zoom · Ziehen: Verschieben · Doppelklick: Ansicht zurücksetzen</div>
    </div>
  </div>

  <div class='footer'>Erstellt mit „FW SAP – Quelldatei Abgleich“</div>
</div>
<script>
const GRAPH_DATA = {graph_json};
let graphInitialized=false;
let graphInitialViewBox=null;
let activeFilter='all';
function setFilter(filter,button){{
  activeFilter=filter;
  document.querySelectorAll('.filter-btn').forEach(b=>b.classList.remove('active'));
  if(button) button.classList.add('active');
  applyFilters();
}}
function applyFilters(){{
  const table=document.getElementById('resultTable');
  if(!table) return;
  const term=(document.getElementById('searchInput').value||'').toLowerCase().trim();
  let visible=0;
  for(const row of table.tBodies[0].rows){{
    const showText=row.innerText.toLowerCase().includes(term);
    const showStatus=activeFilter==='all'||row.dataset.status===activeFilter;
    const show=showText&&showStatus;
    row.style.display=show?'':'none';
    if(show) visible++;
  }}
  document.getElementById('resultCount').textContent=visible+' von '+table.tBodies[0].rows.length;
}}

function showView(view){{
  const table=document.getElementById('tableView');
  const graph=document.getElementById('graphView');
  const tableBtn=document.getElementById('tableViewBtn');
  const graphBtn=document.getElementById('graphViewBtn');
  const showGraph=view==='graph';
  table.classList.toggle('hidden',showGraph);
  graph.classList.toggle('hidden',!showGraph);
  tableBtn.classList.toggle('active',!showGraph);
  graphBtn.classList.toggle('active',showGraph);
  if(showGraph&&!graphInitialized){{ initGraph(); graphInitialized=true; }}
}}
function svgEl(name,attrs){{
  const el=document.createElementNS('http://www.w3.org/2000/svg',name);
  Object.entries(attrs||{{}}).forEach(function(pair){{el.setAttribute(pair[0],String(pair[1]));}});
  return el;
}}
function setViewBox(svg,v){{
  svg.setAttribute('viewBox',v.x+' '+v.y+' '+v.w+' '+v.h);
  svg._vb={{x:v.x,y:v.y,w:v.w,h:v.h}};
}}
function initGraph(){{
  const svg=document.getElementById('graphSvg');
  const edgeLayer=document.getElementById('graphEdges');
  const nodeLayer=document.getElementById('graphNodes');
  if(!svg||!edgeLayer||!nodeLayer) return;
  edgeLayer.innerHTML=''; nodeLayer.innerHTML='';
  const nodes=GRAPH_DATA.nodes.map(function(n){{return Object.assign({{}},n);}});
  const nodeMap=new Map(nodes.map(function(n){{return [n.id,n];}}));
  const root=nodeMap.get('root');
  if(!root) return;
  root.x=0; root.y=0;
  const groups=nodes.filter(function(n){{return n.type==='group';}});
  const groupCount=Math.max(groups.length,1);
  const groupRadius=230;
  groups.forEach(function(g,i){{
    const a=-Math.PI/2+(Math.PI*2*i/groupCount);
    g.angle=a; g.x=Math.cos(a)*groupRadius; g.y=Math.sin(a)*groupRadius;
  }});
  const customersByParent=new Map();
  nodes.filter(function(n){{return n.type==='customer';}}).forEach(function(n){{
    if(!customersByParent.has(n.parent)) customersByParent.set(n.parent,[]);
    customersByParent.get(n.parent).push(n);
  }});
  groups.forEach(function(g){{
    const list=customersByParent.get(g.id)||[];
    const wedge=Math.min(Math.PI*0.78,Math.PI*1.7/groupCount);
    let start=0,ring=0;
    while(start<list.length){{
      const capacity=22+ring*8;
      const current=list.slice(start,start+capacity);
      const radius=105+ring*62;
      current.forEach(function(n,j){{
        const ratio=current.length===1?0.5:j/(current.length-1);
        const a=g.angle-wedge/2+wedge*ratio;
        n.x=g.x+Math.cos(a)*radius;
        n.y=g.y+Math.sin(a)*radius;
      }});
      start+=capacity; ring++;
    }}
  }});
  const positioned=nodes.filter(function(n){{return Number.isFinite(n.x)&&Number.isFinite(n.y);}});
  const xs=positioned.map(function(n){{return n.x;}}),ys=positioned.map(function(n){{return n.y;}});
  const minX=Math.min.apply(null,xs)-150,maxX=Math.max.apply(null,xs)+220,minY=Math.min.apply(null,ys)-150,maxY=Math.max.apply(null,ys)+150;
  graphInitialViewBox={{x:minX,y:minY,w:Math.max(600,maxX-minX),h:Math.max(450,maxY-minY)}};
  setViewBox(svg,graphInitialViewBox);
  GRAPH_DATA.edges.forEach(function(e){{
    const a=nodeMap.get(e.source),b=nodeMap.get(e.target);
    if(!a||!b) return;
    edgeLayer.appendChild(svgEl('line',{{x1:a.x,y1:a.y,x2:b.x,y2:b.y,class:'graph-edge'}}));
  }});
  nodes.forEach(function(n){{
    const g=svgEl('g',{{class:'graph-node '+n.type+' '+n.status,transform:'translate('+n.x+' '+n.y+')'}});
    const radius=n.type==='root'?24:(n.type==='group'?14:6);
    g.appendChild(svgEl('circle',{{r:radius}}));
    const title=svgEl('title',{{}}); title.textContent=n.name||n.label; g.appendChild(title);
    if(n.type!=='customer'){{
      const t=svgEl('text',{{x:radius+8,y:4,class:'graph-label'}}); t.textContent=n.label; g.appendChild(t);
    }}else{{
      const t=svgEl('text',{{x:10,y:4,class:'graph-customer-label'}}); t.textContent=n.label; g.appendChild(t);
    }}
    g.addEventListener('click',function(ev){{ev.stopPropagation();showGraphDetail(n);}});
    nodeLayer.appendChild(g);
  }});
  setupPanZoom(svg);
}}
function setupPanZoom(svg){{
  if(svg._panZoomReady) return; svg._panZoomReady=true;
  let dragging=false,start=null,startVB=null;
  svg.addEventListener('wheel',function(e){{
    e.preventDefault();
    const vb=svg._vb||graphInitialViewBox,rect=svg.getBoundingClientRect();
    const px=(e.clientX-rect.left)/rect.width,py=(e.clientY-rect.top)/rect.height,scale=e.deltaY>0?1.12:0.89;
    const nw=vb.w*scale,nh=vb.h*scale;
    setViewBox(svg,{{x:vb.x+(vb.w-nw)*px,y:vb.y+(vb.h-nh)*py,w:nw,h:nh}});
  }},{{passive:false}});
  svg.addEventListener('pointerdown',function(e){{
    if(e.button!==0) return; dragging=true; start={{x:e.clientX,y:e.clientY}}; startVB=Object.assign({{}},svg._vb||graphInitialViewBox); svg.classList.add('panning'); svg.setPointerCapture(e.pointerId);
  }});
  svg.addEventListener('pointermove',function(e){{
    if(!dragging) return;
    const rect=svg.getBoundingClientRect(),dx=(e.clientX-start.x)*startVB.w/rect.width,dy=(e.clientY-start.y)*startVB.h/rect.height;
    setViewBox(svg,{{x:startVB.x-dx,y:startVB.y-dy,w:startVB.w,h:startVB.h}});
  }});
  svg.addEventListener('pointerup',function(){{dragging=false;svg.classList.remove('panning');}});
  svg.addEventListener('pointercancel',function(){{dragging=false;svg.classList.remove('panning');}});
  svg.addEventListener('dblclick',function(){{if(graphInitialViewBox)setViewBox(svg,graphInitialViewBox);}});
}}
function escGraph(v){{
  return String(v==null?'':v).replace(/[&<>\"]/g,function(c){{return {{'&':'&amp;','<':'&lt;','>':'&gt;','\"':'&quot;'}}[c];}});
}}
function showGraphDetail(n){{
  const box=document.getElementById('graphDetail');
  if(!box) return;
  if(n.type==='root'){{
    const customers=GRAPH_DATA.nodes.filter(function(x){{return x.type==='customer';}}).length;
    const groups=GRAPH_DATA.nodes.filter(function(x){{return x.type==='group';}}).length;
    box.innerHTML="<div class='graph-detail-title'>"+escGraph(n.name)+"</div><div style='color:#aab2c0;font-size:11px'>"+customers+" Kunden · "+groups+" Gruppen</div>";
    return;
  }}
  if(n.type==='group'){{
    const count=GRAPH_DATA.edges.filter(function(e){{return e.source===n.id&&String(e.target).indexOf('customer_')===0;}}).length;
    box.innerHTML="<div class='graph-detail-title'>"+escGraph(n.name)+"</div><div class='graph-detail-row'><span class='graph-detail-key'>Typ</span><span>Blatt</span></div><div class='graph-detail-row'><span class='graph-detail-key'>Kunden</span><span>"+count+"</span></div>";
    return;
  }}
  const colorLabel=n.status==='missing'?'Fehlt in SAP / Liefertag fehlt':(n.status==='extra'?'Zusätzlicher SAP-Tag':'OK');
  box.innerHTML="<div class='graph-detail-title'>"+escGraph(n.name)+"</div>"+
    "<div class='graph-detail-row'><span class='graph-detail-key'>SAP</span><span>"+escGraph(n.sap)+"</span></div>"+
    "<div class='graph-detail-row'><span class='graph-detail-key'>Blatt</span><span>"+escGraph(n.group)+"</span></div>"+
    "<div class='graph-detail-row'><span class='graph-detail-key'>Soll-Tage</span><span>"+escGraph(n.source)+"</span></div>"+
    "<div class='graph-detail-row'><span class='graph-detail-key'>SAP-Tage</span><span>"+escGraph(n.sap_days)+"</span></div>"+
    "<div class='graph-detail-row'><span class='graph-detail-key'>Fehlt</span><span>"+escGraph(n.missing)+"</span></div>"+
    "<div class='graph-detail-row'><span class='graph-detail-key'>Zusätzlich</span><span>"+escGraph(n.extra)+"</span></div>"+
    "<div class='graph-detail-row'><span class='graph-detail-key'>Bewertung</span><span>"+escGraph(colorLabel)+"</span></div>";
}}

applyFilters();
</script>
</body>
</html>"""
    return report.encode("utf-8")


# ---------------------------------------------------------------------------
# Streamlit-Oberfläche
# ---------------------------------------------------------------------------

st.set_page_config(page_title="FW SAP – Quelldatei Abgleich", layout="wide")

st.markdown(
    """
    <style>
        .block-container { padding-top: 2rem; padding-bottom: 2rem; max-width: 1400px; }
        [data-testid="stMetric"] { background: #f7f8fa; border: 1px solid #e3e7ee; border-radius: 14px; padding: 14px 16px; }
        [data-testid="stFileUploader"] section { border: 1px dashed #b9c0cc; border-radius: 14px; background: #fafbfc; }
        div.stButton > button { border-radius: 12px; height: 3rem; font-weight: 700; }
        .hint { color: #586174; font-size: 0.94rem; }
    </style>
    """,
    unsafe_allow_html=True,
)

st.title("FW SAP – Quelldatei Abgleich")
st.markdown(
    """
    <div class="hint">
    <b>Grundlage ist ausschließlich die Quelldatei / Tourenplanung.</b><br>
    Sie ist der Soll-Stand. SAP wird dagegen geprüft. Dabei werden fehlende und zusätzliche SAP-Liefertage sowie komplett fehlende Kunden erkannt.
    Kunden, die nur in SAP stehen, werden ignoriert.
    </div>
    """,
    unsafe_allow_html=True,
)

st.divider()

upload_left, upload_right = st.columns(2)
with upload_left:
    tourenplanung_datei = st.file_uploader(
        "1. Quelldatei / Tourenplanung (Soll)",
        help="Diese Datei hat Vorrang und definiert die richtigen Liefertage.",
        type=["xlsx", "xlsm", "xls"],
        key="tourenplanung_datei",
    )

with upload_right:
    sap_datei = st.file_uploader(
        "2. SAP-Datei (Ist)",
        help="SAP wird gegen die Quelldatei geprüft.",
        type=["xlsx", "xlsm", "xls"],
        key="sap_datei",
    )

# Wichtig: Ergebnisse dürfen nie zu zuvor hochgeladenen Dateien gehören.
# Streamlit behält Session-State über Dateiwechsel hinweg; deshalb wird das
# Ergebnis anhand des tatsächlichen Dateiinhalts eindeutig an beide Uploads gebunden.
current_tour_fingerprint = uploaded_file_fingerprint(tourenplanung_datei)
current_sap_fingerprint = uploaded_file_fingerprint(sap_datei)

existing_result = st.session_state.get("fw_sap_compare_result")
if existing_result:
    result_tour_fp = existing_result.get("tour_fingerprint", "")
    result_sap_fp = existing_result.get("sap_fingerprint", "")
    if (
        result_tour_fp != current_tour_fingerprint
        or result_sap_fp != current_sap_fingerprint
    ):
        st.session_state.pop("fw_sap_compare_result", None)
        existing_result = None
        if tourenplanung_datei is not None or sap_datei is not None:
            st.info("Eine Quelldatei oder SAP-Datei wurde geändert. Bitte den SAP-Abgleich neu erstellen.")

run = st.button("SAP-Abgleich erstellen", type="primary", use_container_width=True)

if run:
    if not sap_datei or not tourenplanung_datei:
        st.error("Bitte Quelldatei und SAP-Datei hochladen.")
        st.stop()

    try:
        tour_df, tour_sheets, missing_tour_sheets, customer_info = read_tourenplanung(tourenplanung_datei)
        days_by_sap, sap_customers, sap_sheet, sap_rows = read_sap_file(sap_datei)

        if tour_df.empty:
            st.error("In der Quelldatei wurden keine gültigen Liefertage erkannt.")
            st.stop()
        if sap_rows == 0:
            st.warning("In der SAP-Datei wurden keine gültigen Liefertage erkannt.")
        if missing_tour_sheets:
            st.warning("Nicht gefundene Quelldatei-Blätter: " + ", ".join(missing_tour_sheets))

        differences = build_sap_differences(tour_df, days_by_sap, sap_customers, customer_info)
        overview = build_customer_overview(tour_df, days_by_sap, sap_customers, customer_info)
        excel_bytes = build_excel(overview, differences)
        html_bytes = build_html_report(overview, differences, sap_sheet, tour_sheets, excel_bytes)

        missing_customers = int((overview["Status"] == "Kunde fehlt in SAP").sum()) if not overview.empty else 0

        st.session_state["fw_sap_compare_result"] = {
            "differences": differences,
            "overview": overview,
            "missing_customers": missing_customers,
            "excel_bytes": excel_bytes,
            "html_bytes": html_bytes,
            "tour_sheets": tour_sheets,
            "sap_sheet": sap_sheet,
            "tour_fingerprint": current_tour_fingerprint,
            "sap_fingerprint": current_sap_fingerprint,
        }

    except Exception as exc:
        import traceback
        st.error(f"Fehler beim Verarbeiten der Dateien: {exc}")
        with st.expander("Technische Details", expanded=False):
            st.code(traceback.format_exc(), language="python")
        st.session_state.pop("fw_sap_compare_result", None)


result = st.session_state.get("fw_sap_compare_result")
if result:
    differences = result["differences"]
    overview = result["overview"]

    st.divider()
    st.subheader("Auswertung erstellt")
    st.caption("Die HTML-Datei enthält die vollständige optische Auswertung und den Excel-Download direkt in der Datei.")

    d1, d2 = st.columns(2)
    with d1:
        st.download_button(
            "HTML-Auswertung herunterladen",
            data=result["html_bytes"],
            file_name="FW_SAP_Quelldatei_Abgleich.html",
            mime="text/html",
            use_container_width=True,
        )
    with d2:
        st.download_button(
            "Excel herunterladen",
            data=result["excel_bytes"],
            file_name="FW_SAP_Quelldatei_Abgleich.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )

    m1, m2, m3 = st.columns(3)
    m1.metric("Geprüfte Kunden", len(overview))
    m2.metric("SAP-Abweichungen", len(differences))
    m3.metric("Kunden fehlen in SAP", result.get("missing_customers", 0))

    if differences.empty:
        st.success("Keine Abweichungen: SAP stimmt mit der Quelldatei überein.")
    else:
        st.warning(f"{len(differences)} Kunde(n) mit SAP-Abweichung gefunden. Für die vollständige Übersicht bitte die HTML-Auswertung öffnen.")
